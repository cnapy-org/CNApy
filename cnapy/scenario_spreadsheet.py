"""Reading and writing CNApy scenarios as spreadsheets (.xlsx, works with Excel and LibreOffice).

The workbook has the sheets General, Fluxes, Reactions, Constraints, Annotations and Help.
In each data sheet row 1 is a header, rows whose first cell is empty or starts with '#'
are skipped and missing sheets are treated as empty. read_scenario_xlsx returns a dictionary
in the same format as the JSON content of a .scen file so that it can be loaded the same way.
"""
from math import isinf
from typing import Dict, List

import cobra
import openpyxl
from openpyxl.styles import Font, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from straindesign.parse_constr import linexpr2dict, linexprdict2str, lineq2list

GENERAL = "General"
FLUXES = "Fluxes"
REACTIONS = "Reactions"
CONSTRAINTS = "Constraints"
ANNOTATIONS = "Annotations"
HELP = "Help"

DESCRIPTION = "Description"
OBJECTIVE = "Objective"
DIRECTION = "Direction"
USE_OBJECTIVE = "Use scenario objective"

HEADERS = {
    GENERAL: ["Setting", "Value"],
    FLUXES: ["Reaction ID", "Lower bound", "Upper bound", "Pinned"],
    REACTIONS: ["Reaction ID", "Equation", "Lower bound", "Upper bound"],
    CONSTRAINTS: ["Constraint"],
    ANNOTATIONS: ["Reaction ID", "Key", "Value"],
}
COLUMN_WIDTHS = {
    GENERAL: [24, 80],
    FLUXES: [24, 14, 14, 10],
    REACTIONS: [24, 80, 14, 14],
    CONSTRAINTS: [80],
    ANNOTATIONS: [24, 24, 40],
}

EXAMPLES = {
    FLUXES: [["# EX_glc__D_e", -10, -10, "x"], ["# EX_o2_e", -5, 0, ""], ["# BIOMASS", "", "", "x"]],
    REACTIONS: [["# EX_prod", "prod_c -->", 0, 1000], ["# PROD", "2 pyr_c + nadh_c --> prod_c + nad_c", "", ""]],
    CONSTRAINTS: [["# EX_prod + 0.5 EX_glc__D_e >= 0"]],
    ANNOTATIONS: [["# PROD", "ec-code", "1.1.1.1"]],
}

HELP_TEXT = """\
How to use this CNApy scenario spreadsheet

Load it in CNApy with Scenario > Load scenario... (choose the .xlsx file).
When editing in LibreOffice keep the .xlsx format when saving.

General rules
- Row 1 of each sheet is a header and is not read.
- Rows whose first cell is empty or starts with # are ignored (use # for comments or to disable a row).
- Sheets that you do not need can be left empty or deleted. This Help sheet is not read.
- Bounds may be numbers or inf / -inf.

Fluxes: scenario flux values of model reactions.
- Give lower and upper bound; if only one of them is given the flux is fixed to that value.
- Pinned: x (or yes/true/1) pins the reaction to the top of the reaction list.
- A row with a reaction ID, no bounds and a mark in Pinned only pins the reaction.

General: settings in column B, identified by the name in column A.
- Description: free text describing the scenario.
- Objective: linear expression of reaction IDs, e.g. BIOMASS or 2 R1 - 0.5 R3.
- Direction: max or min.
- Use scenario objective: yes or no; when empty, yes if an objective is given.

Reactions: additional reactions that are only part of this scenario.
- Equation in COBRApy notation, e.g. 2 pyr_c + nadh_c --> prod_c + nad_c (also <=> and <--).
- Unknown metabolites are created when the scenario is applied.
- If the bounds are left empty they are derived from the arrow in the equation.

Constraints: one linear (in)equality per row, e.g. EX_prod + 0.5 EX_glc__D_e >= 0 (also <= and =).

Annotations: reaction annotations (key and value) that are set while the scenario is active.
"""

_TRUE_STRINGS = {"x", "yes", "y", "true", "1"}
_FALSE_STRINGS = {"no", "n", "false", "0"}


class ScenarioSpreadsheetError(Exception):
    def __init__(self, messages: List[str]):
        super().__init__("\n".join(messages))
        self.messages = messages


def _number_cell(value: float):
    # infinite values cannot be stored as numbers in xlsx files
    if isinf(value):
        return "inf" if value > 0 else "-inf"
    return value


def _add_sheet(wb: openpyxl.Workbook, title: str, rows: List[List]):
    ws = wb.create_sheet(title)
    ws.append(HEADERS[title])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    for i, width in enumerate(COLUMN_WIDTHS[title], start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for row in rows:
        ws.append(row)
    return ws


def write_scenario_xlsx(scenario, filename: str, include_examples: bool = False):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    rows = [[reac_id, _number_cell(lb), _number_cell(ub), "x" if reac_id in scenario.pinned_reactions else ""]
            for reac_id, (lb, ub) in scenario.items()]
    rows += [[reac_id, "", "", "x"] for reac_id in sorted(scenario.pinned_reactions) if reac_id not in scenario]
    _add_sheet(wb, FLUXES, (EXAMPLES[FLUXES] if include_examples else []) + rows)

    ws = _add_sheet(wb, GENERAL, [
        [DESCRIPTION, scenario.description],
        [OBJECTIVE, linexprdict2str(scenario.objective_coefficients)],
        [DIRECTION, scenario.objective_direction],
        [USE_OBJECTIVE, "yes" if scenario.use_scenario_objective else "no"],
        [],
        ["# Instructions can be found in the Help sheet."]])
    ws["B2"].alignment = Alignment(wrap_text=True, vertical="top")
    for cell, choices in (("B4", '"max,min"'), ("B5", '"yes,no"')):
        dv = DataValidation(type="list", formula1=choices, allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(cell)

    rows = []
    for reac_id, (metabolites, lb, ub) in scenario.reactions.items():
        reaction = cobra.Reaction(reac_id, lower_bound=lb, upper_bound=ub)
        reaction.add_metabolites({cobra.Metabolite(met_id): coeff for met_id, coeff in metabolites.items()})
        rows.append([reac_id, reaction.build_reaction_string(), _number_cell(lb), _number_cell(ub)])
    _add_sheet(wb, REACTIONS, (EXAMPLES[REACTIONS] if include_examples else []) + rows)

    rows = [[linexprdict2str(constr[0]) + " " + constr[1] + " " + str(constr[2])]
            for constr in scenario.constraints if constr[0] is not None]
    _add_sheet(wb, CONSTRAINTS, (EXAMPLES[CONSTRAINTS] if include_examples else []) + rows)

    rows = [[a["reaction_id"], a["key"], a["value"]] for a in scenario.annotations]
    _add_sheet(wb, ANNOTATIONS, (EXAMPLES[ANNOTATIONS] if include_examples else []) + rows)

    ws = wb.create_sheet(HELP)
    ws.column_dimensions["A"].width = 120
    for line in HELP_TEXT.splitlines():
        ws.append([line])
    ws["A1"].font = Font(bold=True, size=14)

    wb.active = 0
    wb.save(filename)


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


class _SheetReader:
    """Collects the error messages of all sheets so that they can be reported together."""

    def __init__(self, filename: str):
        self.wb = openpyxl.load_workbook(filename, read_only=True, data_only=True)
        self.sheets = {name.casefold(): name for name in self.wb.sheetnames}
        self.errors: List[str] = []
        self.sheet = ""

    def rows(self, title: str):
        """yields (row number, row values) of the data rows, padded to the header length"""
        self.sheet = title
        name = self.sheets.get(title.casefold())
        if name is None:
            return
        num_columns = len(HEADERS[title])
        for row_number, row in enumerate(self.wb[name].iter_rows(min_row=2, values_only=True), start=2):
            row = list(row[:num_columns]) + [None]*(num_columns - len(row))
            first = _text(row[0])
            if len(first) > 0 and not first.startswith("#"):
                yield row_number, row

    def error(self, row_number: int, column: int, message: str):
        self.errors.append(f"{self.sheet}!{get_column_letter(column + 1)}{row_number}: {message}")

    def number(self, row_number: int, row: List, column: int):
        value = row[column]
        if value is None or isinstance(value, bool):
            if value is not None:
                self.error(row_number, column, f"'{value}' is not a number")
            return None
        if isinstance(value, (int, float)):
            return float(value)
        text = _text(value)
        if len(text) == 0:
            return None
        try:
            return float(text.replace(",", "."))
        except ValueError:
            self.error(row_number, column, f"'{text}' is not a number")
            return None

    def flag(self, row_number: int, row: List, column: int, default=False) -> bool:
        value = row[column]
        if isinstance(value, bool):
            return value
        text = _text(value).casefold()
        if len(text) == 0:
            return default
        if text in _TRUE_STRINGS:
            return True
        if text not in _FALSE_STRINGS:
            self.error(row_number, column, f"'{_text(value)}' is not understood, use yes or no")
        return False


def read_scenario_xlsx(filename: str, model_reaction_ids) -> Dict:
    """Reads a scenario spreadsheet and returns it in the format of a .scen JSON dictionary.
    model_reaction_ids are needed to parse the objective and the constraints; all errors
    found are collectively raised as ScenarioSpreadsheetError."""
    reader = _SheetReader(filename)

    fluxes = {}
    pinned_reactions = []
    for row_number, row in reader.rows(FLUXES):
        reac_id = _text(row[0])
        num_errors = len(reader.errors)
        lb = reader.number(row_number, row, 1)
        ub = reader.number(row_number, row, 2)
        pinned = reader.flag(row_number, row, 3)
        if lb is None and ub is None:
            if not pinned and len(reader.errors) == num_errors:
                reader.error(row_number, 1, f"no flux value given for {reac_id}")
        else:
            if reac_id in fluxes:
                reader.error(row_number, 0, f"flux value for {reac_id} is given more than once")
            if lb is None:
                lb = ub
            elif ub is None:
                ub = lb
            if lb > ub:
                reader.error(row_number, 1, f"lower bound {lb} is greater than upper bound {ub}")
            fluxes[reac_id] = (lb, ub)
        if pinned:
            pinned_reactions.append(reac_id)

    reactions = {}
    model = cobra.Model() # only used to parse the equations
    for row_number, row in reader.rows(REACTIONS):
        reac_id = _text(row[0])
        if reac_id in reactions:
            reader.error(row_number, 0, f"reaction {reac_id} is defined more than once")
            continue
        reaction = cobra.Reaction(reac_id)
        model.add_reactions([reaction])
        try:
            reaction.build_reaction_from_string(_text(row[1]), verbose=False)
        except Exception as e:
            reader.error(row_number, 1, f"cannot parse equation '{_text(row[1])}' ({e})")
        lb = reader.number(row_number, row, 2)
        ub = reader.number(row_number, row, 3)
        reactions[reac_id] = [{m.id: c for m, c in reaction.metabolites.items()},
                              reaction.lower_bound if lb is None else lb,
                              reaction.upper_bound if ub is None else ub]

    reaction_ids = list(model_reaction_ids) + [r for r in reactions if r not in model_reaction_ids]

    constraints = []
    for row_number, row in reader.rows(CONSTRAINTS):
        try:
            constraints.append(list(lineq2list([_text(row[0])], reaction_ids)[0]))
        except Exception as e:
            reader.error(row_number, 0, f"cannot parse constraint '{_text(row[0])}' ({e})")

    annotations = []
    for row_number, row in reader.rows(ANNOTATIONS):
        key = _text(row[1])
        if len(key) == 0:
            reader.error(row_number, 1, "annotation key is missing")
        annotations.append({"reaction_id": _text(row[0]), "key": key, "value": _text(row[2])})

    general = {}
    for row_number, row in reader.rows(GENERAL):
        general[_text(row[0]).casefold()] = (row_number, row)
    description = ""
    if DESCRIPTION.casefold() in general:
        value = general[DESCRIPTION.casefold()][1][1]
        description = "" if value is None else str(value)
    objective_coefficients = {}
    if OBJECTIVE.casefold() in general:
        row_number, row = general[OBJECTIVE.casefold()]
        try:
            objective_coefficients = linexpr2dict(_text(row[1]), reaction_ids)
        except Exception as e:
            reader.error(row_number, 1, f"cannot parse objective '{_text(row[1])}' ({e})")
    objective_direction = "max"
    if DIRECTION.casefold() in general:
        row_number, row = general[DIRECTION.casefold()]
        direction = _text(row[1]).casefold()
        if direction in ("min", "minimize", "minimise"):
            objective_direction = "min"
        elif direction not in ("", "max", "maximize", "maximise"):
            reader.error(row_number, 1, f"'{_text(row[1])}' is not a valid direction, use max or min")
    use_scenario_objective = len(objective_coefficients) > 0
    if USE_OBJECTIVE.casefold() in general:
        row_number, row = general[USE_OBJECTIVE.casefold()]
        use_scenario_objective = reader.flag(row_number, row, 1, default=use_scenario_objective)

    reader.wb.close()
    if len(reader.errors) > 0:
        raise ScenarioSpreadsheetError(reader.errors)

    return {'fluxes': fluxes, 'pinned_reactions': pinned_reactions, 'description': description,
            'objective_direction': objective_direction, 'objective_coefficients': objective_coefficients,
            'use_scenario_objective': use_scenario_objective, 'reactions': reactions,
            'constraints': constraints, 'annotations': annotations, 'version': 3}
